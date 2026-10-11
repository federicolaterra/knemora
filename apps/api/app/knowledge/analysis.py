"""Deterministic, bounded analysis for immutable CSV and XLSX sources."""

from __future__ import annotations

import hashlib
import io
import uuid
from enum import StrEnum
from typing import Any, Literal

import polars as pl
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import Field, model_validator
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.db.models import DatasetRecord, Document, DocumentVersion
from app.ingestion.storage import SourceStorage
from app.knowledge.errors import DatasetNotFoundError, InvalidDatasetAnalysisError
from app.knowledge.types import StrictModel


class FilterOperator(StrEnum):
    EQ = "eq"
    NE = "ne"
    GT = "gt"
    GTE = "gte"
    LT = "lt"
    LTE = "lte"
    CONTAINS = "contains"
    IN = "in"
    IS_NULL = "is_null"
    NOT_NULL = "not_null"


class AnalysisMetric(StrEnum):
    SUM = "sum"
    AVERAGE = "average"
    MEDIAN = "median"
    COUNT = "count"


class DatasetFilter(StrictModel):
    column: str = Field(min_length=1, max_length=300)
    operator: FilterOperator
    value: Any = None


class DatasetAnalysisRequest(StrictModel):
    operation: Literal["aggregate", "describe", "trend", "comparison", "top", "bottom"]
    sheet: str | None = Field(default=None, max_length=300)
    filters: tuple[DatasetFilter, ...] = ()
    group_by: tuple[str, ...] = ()
    value_column: str | None = Field(default=None, max_length=300)
    metrics: tuple[AnalysisMetric, ...] = ()
    order_by: str | None = Field(default=None, max_length=300)
    limit: int = Field(default=50, ge=1, le=5_000)
    offset: int = Field(default=0, ge=0)
    chart_type: Literal["bar", "line"] | None = None

    @model_validator(mode="after")
    def validate_operation(self) -> DatasetAnalysisRequest:
        numeric_metrics = {AnalysisMetric.SUM, AnalysisMetric.AVERAGE, AnalysisMetric.MEDIAN}
        if numeric_metrics.intersection(self.metrics) and not self.value_column:
            raise ValueError("sum, average, and median require value_column")
        if self.operation in {"trend", "comparison"} and not self.group_by:
            raise ValueError(f"{self.operation} requires at least one group_by column")
        if self.operation in {"top", "bottom"} and not (self.group_by or self.value_column):
            raise ValueError(f"{self.operation} requires group_by or value_column")
        return self


class DatasetProvenance(StrictModel):
    document_id: str
    document_version_id: str
    original_filename: str
    sha256: str
    immutable_source_uri: str
    sheet: str | None
    source_columns: tuple[str, ...]
    source_row_count: int
    filtered_row_count: int
    filters: tuple[DatasetFilter, ...]


class ChartSeries(StrictModel):
    name: str
    values: tuple[Any, ...]


class ChartData(StrictModel):
    type: Literal["bar", "line"]
    x_column: str
    x: tuple[Any, ...]
    series: tuple[ChartSeries, ...]


class DatasetAnalysisResult(StrictModel):
    operation: str
    columns: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]
    row_count: int
    truncated: bool
    chart: ChartData | None
    provenance: DatasetProvenance


class DatasetAnalysisService:
    """Run an allow-listed Polars calculation over one owner-scoped source."""

    def __init__(self, *, session: Session, settings: Settings, owner_id: uuid.UUID) -> None:
        self.session = session
        self.settings = settings
        self.owner_id = owner_id
        self.storage = SourceStorage(settings.source_store_path)

    def analyze(
        self, document_version_id: uuid.UUID, request: DatasetAnalysisRequest
    ) -> DatasetAnalysisResult:
        dataset, version, document = self._load_dataset(document_version_id)
        frame, sheet = self._read_frame(version, dataset, request.sheet)
        if frame.height > self.settings.dataset_analysis_max_rows:
            raise InvalidDatasetAnalysisError(
                f"Dataset has {frame.height} rows; the local analysis limit is "
                f"{self.settings.dataset_analysis_max_rows}."
            )
        filtered = self._apply_filters(frame, request.filters)
        result = self._calculate(filtered, request)
        maximum = min(request.limit, self.settings.dataset_analysis_max_result_rows)
        row_count = result.height
        truncated = row_count > maximum
        result = result.slice(request.offset, maximum)
        rows = tuple(self._json_row(row) for row in result.to_dicts())
        chart = self._chart(result, request)
        used_columns = tuple(
            dict.fromkeys(
                [
                    *(item.column for item in request.filters),
                    *request.group_by,
                    *([request.value_column] if request.value_column else []),
                    *([request.order_by] if request.order_by else []),
                ]
            )
        )
        if not used_columns:
            used_columns = tuple(frame.columns)
        return DatasetAnalysisResult(
            operation=request.operation,
            columns=tuple(result.columns),
            rows=rows,
            row_count=row_count,
            truncated=truncated,
            chart=chart,
            provenance=DatasetProvenance(
                document_id=str(document.id),
                document_version_id=str(version.id),
                original_filename=document.original_filename,
                sha256=version.sha256,
                immutable_source_uri=version.storage_uri,
                sheet=sheet,
                source_columns=used_columns,
                source_row_count=frame.height,
                filtered_row_count=filtered.height,
                filters=request.filters,
            ),
        )

    def _load_dataset(
        self, version_id: uuid.UUID
    ) -> tuple[DatasetRecord, DocumentVersion, Document]:
        row = self.session.execute(
            select(DatasetRecord, DocumentVersion, Document)
            .join(
                DocumentVersion,
                DatasetRecord.document_version_id == DocumentVersion.id,
            )
            .join(Document, DocumentVersion.document_id == Document.id)
            .where(
                DocumentVersion.id == version_id,
                Document.owner_id == self.owner_id,
                Document.deleted_at.is_(None),
            )
        ).one_or_none()
        if row is None:
            raise DatasetNotFoundError(str(version_id))
        return row

    def _read_frame(
        self,
        version: DocumentVersion,
        dataset: DatasetRecord,
        requested_sheet: str | None,
    ) -> tuple[pl.DataFrame, str | None]:
        path = self.storage.resolve_relative(version.storage_uri)
        data = self.storage.read_original(path)
        if hashlib.sha256(data).hexdigest() != version.sha256:
            raise InvalidDatasetAnalysisError("The immutable dataset failed checksum verification.")
        kind = str(dataset.schema_metadata.get("kind", ""))
        try:
            if kind == "csv":
                if requested_sheet is not None:
                    raise InvalidDatasetAnalysisError("CSV datasets do not have sheets.")
                delimiter = str(dataset.schema_metadata.get("delimiter", ","))
                frame = pl.read_csv(
                    io.BytesIO(data),
                    separator=delimiter,
                    try_parse_dates=True,
                    infer_schema_length=10_000,
                )
                return frame, None
            if kind == "xlsx":
                return self._read_xlsx(data, dataset.schema_metadata, requested_sheet)
        except InvalidDatasetAnalysisError:
            raise
        except Exception as exc:
            raise InvalidDatasetAnalysisError(
                f"Unable to read the structured dataset: {type(exc).__name__}."
            ) from exc
        raise InvalidDatasetAnalysisError(f"Unsupported structured dataset kind: {kind!r}.")

    @staticmethod
    def _read_xlsx(
        data: bytes, metadata: dict[str, Any], requested_sheet: str | None
    ) -> tuple[pl.DataFrame, str]:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            known_sheets = [str(item["name"]) for item in metadata.get("sheets", [])]
            sheet = requested_sheet or (known_sheets[0] if known_sheets else workbook.sheetnames[0])
            if sheet not in workbook.sheetnames:
                raise InvalidDatasetAnalysisError(f"Workbook sheet {sheet!r} was not found.")
            rows = list(workbook[sheet].iter_rows(values_only=True))
            while rows and all(value is None for value in rows[-1]):
                rows.pop()
            if not rows:
                raise InvalidDatasetAnalysisError(f"Workbook sheet {sheet!r} is empty.")
            headers = [str(value).strip() if value is not None else "" for value in rows[0]]
            if any(not header for header in headers) or len(headers) != len(set(headers)):
                raise InvalidDatasetAnalysisError(
                    "Dataset columns must have unique, non-blank header names."
                )
            frame = pl.DataFrame(rows[1:], schema=headers, orient="row", strict=False)
            return frame, sheet
        finally:
            workbook.close()

    def _apply_filters(
        self, frame: pl.DataFrame, filters: tuple[DatasetFilter, ...]
    ) -> pl.DataFrame:
        result = frame
        for item in filters:
            self._require_columns(result, (item.column,))
            column = pl.col(item.column)
            if item.operator == FilterOperator.EQ:
                expression = column == item.value
            elif item.operator == FilterOperator.NE:
                expression = column != item.value
            elif item.operator == FilterOperator.GT:
                expression = column > item.value
            elif item.operator == FilterOperator.GTE:
                expression = column >= item.value
            elif item.operator == FilterOperator.LT:
                expression = column < item.value
            elif item.operator == FilterOperator.LTE:
                expression = column <= item.value
            elif item.operator == FilterOperator.CONTAINS:
                expression = column.cast(pl.String).str.contains(str(item.value), literal=True)
            elif item.operator == FilterOperator.IN:
                if not isinstance(item.value, list):
                    raise InvalidDatasetAnalysisError("The 'in' filter requires a JSON list.")
                expression = column.is_in(item.value)
            elif item.operator == FilterOperator.IS_NULL:
                expression = column.is_null()
            else:
                expression = column.is_not_null()
            try:
                result = result.filter(expression)
            except Exception as exc:
                raise InvalidDatasetAnalysisError(
                    f"Filter on column {item.column!r} is incompatible with its data type."
                ) from exc
        return result

    def _calculate(self, frame: pl.DataFrame, request: DatasetAnalysisRequest) -> pl.DataFrame:
        self._require_columns(
            frame,
            (
                *request.group_by,
                *([request.value_column] if request.value_column else []),
            ),
        )
        if request.operation == "describe":
            return self._describe(frame, request.value_column)
        result = self._aggregate(frame, request)
        if request.operation in {"top", "bottom"}:
            sort_column = request.order_by or self._primary_value_column(result, request)
            self._require_columns(result, (sort_column,))
            result = result.sort(sort_column, descending=request.operation == "top")
        elif request.operation == "trend":
            result = result.sort(list(request.group_by))
        elif request.order_by:
            self._require_columns(result, (request.order_by,))
            result = result.sort(request.order_by)
        return result

    def _aggregate(self, frame: pl.DataFrame, request: DatasetAnalysisRequest) -> pl.DataFrame:
        metrics = request.metrics or (AnalysisMetric.COUNT,)
        expressions: list[pl.Expr] = []
        for metric in metrics:
            if metric == AnalysisMetric.COUNT:
                expressions.append(pl.len().alias("count"))
                continue
            if request.value_column is None:
                raise InvalidDatasetAnalysisError(f"{metric.value} requires value_column.")
            column = pl.col(request.value_column)
            if metric == AnalysisMetric.SUM:
                expressions.append(column.sum().alias(f"{request.value_column}_sum"))
            elif metric == AnalysisMetric.AVERAGE:
                expressions.append(column.mean().alias(f"{request.value_column}_average"))
            elif metric == AnalysisMetric.MEDIAN:
                expressions.append(column.median().alias(f"{request.value_column}_median"))
        try:
            if request.group_by:
                return frame.group_by(list(request.group_by), maintain_order=True).agg(expressions)
            return frame.select(expressions)
        except Exception as exc:
            raise InvalidDatasetAnalysisError(
                "The requested aggregation is incompatible with the selected columns."
            ) from exc

    def _describe(self, frame: pl.DataFrame, value_column: str | None) -> pl.DataFrame:
        columns = (
            [value_column]
            if value_column
            else [name for name, dtype in frame.schema.items() if dtype.is_numeric()]
        )
        if not columns:
            raise InvalidDatasetAnalysisError("No numeric columns are available to describe.")
        self._require_columns(frame, tuple(columns))
        rows: list[dict[str, Any]] = []
        for name in columns:
            series = frame.get_column(name)
            if not series.dtype.is_numeric():
                raise InvalidDatasetAnalysisError(f"Column {name!r} is not numeric.")
            rows.append(
                {
                    "column": name,
                    "count": series.len() - series.null_count(),
                    "null_count": series.null_count(),
                    "mean": series.mean(),
                    "std": series.std(),
                    "min": series.min(),
                    "median": series.median(),
                    "max": series.max(),
                }
            )
        return pl.DataFrame(rows)

    @staticmethod
    def _primary_value_column(result: pl.DataFrame, request: DatasetAnalysisRequest) -> str:
        candidates = [column for column in result.columns if column not in request.group_by]
        if not candidates:
            raise InvalidDatasetAnalysisError("No computed value is available for ranking.")
        return candidates[0]

    @staticmethod
    def _chart(result: pl.DataFrame, request: DatasetAnalysisRequest) -> ChartData | None:
        if request.chart_type is None or result.is_empty():
            return None
        x_column = request.group_by[0] if request.group_by else result.columns[0]
        value_columns = [
            name for name, dtype in result.schema.items() if name != x_column and dtype.is_numeric()
        ]
        if not value_columns:
            raise InvalidDatasetAnalysisError("Chart output requires a computed numeric series.")
        return ChartData(
            type=request.chart_type,
            x_column=x_column,
            x=tuple(DatasetAnalysisService._json_value(value) for value in result[x_column]),
            series=tuple(
                ChartSeries(
                    name=name,
                    values=tuple(
                        DatasetAnalysisService._json_value(value) for value in result[name]
                    ),
                )
                for name in value_columns
            ),
        )

    @staticmethod
    def _require_columns(frame: pl.DataFrame, columns: tuple[str, ...]) -> None:
        missing = sorted(set(columns) - set(frame.columns))
        if missing:
            raise InvalidDatasetAnalysisError(f"Unknown dataset column(s): {', '.join(missing)}.")

    @staticmethod
    def _json_value(value: Any) -> Any:
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return value

    @staticmethod
    def _json_row(row: dict[str, Any]) -> dict[str, Any]:
        return {key: DatasetAnalysisService._json_value(value) for key, value in row.items()}


__all__ = [
    "AnalysisMetric",
    "ChartData",
    "DatasetAnalysisRequest",
    "DatasetAnalysisResult",
    "DatasetAnalysisService",
    "DatasetFilter",
    "FilterOperator",
]
