"""Bounded parser implementations for the Milestone 03 source allow-list."""

from __future__ import annotations

import csv
import io
import re
from importlib.metadata import PackageNotFoundError, version
from typing import Any

from docx import Document as OpenDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader

from app.ingestion.errors import ParserFailure
from app.ingestion.types import DocumentParser, ParsedDocument, SourceBlock, ValidatedUpload

_MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_TABULAR_BLOCK_ROWS = 100


def _package_version(distribution: str) -> str:
    try:
        return version(distribution)
    except PackageNotFoundError:
        return "unknown"


def _markdown_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _markdown_table(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [None] * (width - len(row)) for row in rows]
    header = padded[0]
    body = padded[1:]
    lines = [
        "| " + " | ".join(_markdown_cell(value) for value in header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(_markdown_cell(value) for value in row) + " |" for row in body)
    return "\n".join(lines)


class TextParser:
    name = "knemora-text"
    version = "1.0.0"
    extensions = frozenset({".txt"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        text = upload.data.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        line_count = max(1, len(text.splitlines()))
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=(SourceBlock(text.rstrip(), {"line_start": 1, "line_end": line_count}),),
            document_metadata={"line_count": line_count},
        )


class MarkdownParser:
    name = "knemora-markdown"
    version = "1.0.0"
    extensions = frozenset({".md", ".markdown"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        text = upload.data.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        lines = text.splitlines()
        headings: list[str] = []
        blocks: list[SourceBlock] = []
        block_lines: list[str] = []
        block_start = 1
        block_heading_path: list[str] = []

        def flush(end_line: int) -> None:
            content = "\n".join(block_lines).strip()
            if content:
                locator: dict[str, Any] = {"line_start": block_start, "line_end": end_line}
                if block_heading_path:
                    locator["heading_path"] = block_heading_path.copy()
                blocks.append(SourceBlock(content, locator))

        for line_number, line in enumerate(lines, start=1):
            match = _MARKDOWN_HEADING.match(line)
            if match:
                flush(line_number - 1)
                level = len(match.group(1))
                title = match.group(2).strip()
                headings[level - 1 :] = [title]
                block_lines = [line]
                block_start = line_number
                block_heading_path = headings.copy()
            else:
                block_lines.append(line)
        flush(max(1, len(lines)))
        if not blocks:
            blocks.append(SourceBlock("", {"line_start": 1, "line_end": 1}))
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=tuple(blocks),
            document_metadata={"line_count": max(1, len(lines))},
        )


class CsvParser:
    name = "python-csv"
    version = "stdlib"
    extensions = frozenset({".csv"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        text = upload.data.decode("utf-8-sig")
        try:
            sample = text[:8192]
            try:
                dialect = (
                    csv.Sniffer().sniff(sample, delimiters=",;\t|") if sample.strip() else csv.excel
                )
            except csv.Error:
                dialect = csv.excel
            rows = [list(row) for row in csv.reader(io.StringIO(text), dialect, strict=True)]
        except csv.Error as exc:
            raise ParserFailure(f"Malformed CSV: {exc}") from exc
        if not rows:
            raise ParserFailure("The CSV contains no rows.")
        width = max(len(row) for row in rows)
        if width == 0:
            raise ParserFailure("The CSV contains no columns.")
        header = rows[0]
        blocks: list[SourceBlock] = []
        data_rows = rows[1:]
        if not data_rows:
            blocks.append(SourceBlock(_markdown_table([header]), {"row_start": 1, "row_end": 1}))
        for offset in range(0, len(data_rows), _TABULAR_BLOCK_ROWS):
            chunk = data_rows[offset : offset + _TABULAR_BLOCK_ROWS]
            row_start = offset + 2
            blocks.append(
                SourceBlock(
                    _markdown_table([header, *chunk]),
                    {"row_start": row_start, "row_end": row_start + len(chunk) - 1},
                )
            )
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=tuple(blocks),
            document_metadata={"row_count": len(rows), "column_count": width},
            dataset_metadata={
                "kind": "csv",
                "columns": [str(value) for value in header],
                "row_count": max(0, len(rows) - 1),
                "column_count": width,
                "delimiter": dialect.delimiter,
            },
        )


class PdfParser:
    name = "pypdf"
    version = _package_version("pypdf")
    extensions = frozenset({".pdf"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        try:
            reader = PdfReader(io.BytesIO(upload.data), strict=True)
            if reader.is_encrypted:
                raise ParserFailure("Encrypted PDFs are not supported.")
            blocks = tuple(
                SourceBlock((page.extract_text() or "").strip(), {"page": page_number})
                for page_number, page in enumerate(reader.pages, start=1)
            )
        except ParserFailure:
            raise
        except Exception as exc:
            raise ParserFailure(f"Malformed PDF: {type(exc).__name__}") from exc
        if not blocks:
            raise ParserFailure("The PDF contains no pages.")
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=blocks,
            document_metadata={"page_count": len(blocks)},
        )


class DocxParser:
    name = "python-docx"
    version = _package_version("python-docx")
    extensions = frozenset({".docx"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        try:
            document = OpenDocument(io.BytesIO(upload.data))
            headings: list[str] = []
            blocks: list[SourceBlock] = []
            paragraph_number = 0
            table_number = 0
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    paragraph_number += 1
                    text = item.text.strip()
                    style_name = item.style.name if item.style is not None else ""
                    match = re.fullmatch(r"Heading ([1-6])", style_name)
                    if match and text:
                        level = int(match.group(1))
                        headings[level - 1 :] = [text]
                    if text:
                        locator: dict[str, Any] = {"paragraph": paragraph_number}
                        if headings:
                            locator["section"] = headings.copy()
                        blocks.append(SourceBlock(text, locator))
                elif isinstance(item, Table):
                    table_number += 1
                    rows = [[cell.text for cell in row.cells] for row in item.rows]
                    if rows:
                        locator = {"table": table_number, "row_start": 1, "row_end": len(rows)}
                        if headings:
                            locator["section"] = headings.copy()
                        blocks.append(SourceBlock(_markdown_table(rows), locator))
        except Exception as exc:
            raise ParserFailure(f"Malformed DOCX: {type(exc).__name__}") from exc
        if not blocks:
            raise ParserFailure("The DOCX contains no extractable content.")
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=tuple(blocks),
            document_metadata={
                "paragraph_count": paragraph_number,
                "table_count": table_number,
            },
        )


class XlsxParser:
    name = "openpyxl"
    version = _package_version("openpyxl")
    extensions = frozenset({".xlsx"})

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        try:
            workbook = load_workbook(
                io.BytesIO(upload.data), read_only=True, data_only=True, keep_links=False
            )
            blocks: list[SourceBlock] = []
            sheets: list[dict[str, Any]] = []
            for worksheet in workbook.worksheets:
                rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
                while rows and all(value is None for value in rows[-1]):
                    rows.pop()
                width = max((len(row) for row in rows), default=0)
                sheets.append(
                    {"name": worksheet.title, "row_count": len(rows), "column_count": width}
                )
                for offset in range(0, len(rows), _TABULAR_BLOCK_ROWS):
                    chunk = rows[offset : offset + _TABULAR_BLOCK_ROWS]
                    blocks.append(
                        SourceBlock(
                            f"## Sheet: {worksheet.title}\n\n{_markdown_table(chunk)}",
                            {
                                "sheet": worksheet.title,
                                "row_start": offset + 1,
                                "row_end": offset + len(chunk),
                            },
                        )
                    )
            workbook.close()
        except Exception as exc:
            raise ParserFailure(f"Malformed XLSX: {type(exc).__name__}") from exc
        if not blocks:
            raise ParserFailure("The XLSX contains no populated cells.")
        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            blocks=tuple(blocks),
            document_metadata={"sheet_count": len(sheets)},
            dataset_metadata={"kind": "xlsx", "sheets": sheets},
        )


class ParserRegistry:
    """Explicit parser allow-list; new formats must register deliberately."""

    def __init__(self, parsers: tuple[DocumentParser, ...] | None = None) -> None:
        configured = parsers or (
            PdfParser(),
            TextParser(),
            MarkdownParser(),
            DocxParser(),
            CsvParser(),
            XlsxParser(),
        )
        self._by_extension = {
            extension: parser for parser in configured for extension in parser.extensions
        }

    def parse(self, upload: ValidatedUpload) -> ParsedDocument:
        parser = self._by_extension.get(upload.extension)
        if parser is None:
            raise ParserFailure(f"No parser is registered for {upload.extension}.")
        return parser.parse(upload)
