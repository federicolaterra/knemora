"""Regenerate the deterministic binary ingestion fixtures."""

from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook  # type: ignore[import-untyped]

FIXTURE_DIRECTORY = Path(__file__).parent


def build_pdf() -> bytes:
    content = b"BT /F1 12 Tf 72 720 Td (Knemora PDF fixture) Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"endstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(f"{number} 0 obj\n".encode())
        result.extend(body)
        result.extend(b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(result)


def write_docx() -> None:
    document = Document()
    document.add_heading("Quality System", level=1)
    document.add_paragraph("Controlled DOCX fixture content.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Batch"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "B-001"
    table.cell(1, 1).text = "Released"
    document.save(str(FIXTURE_DIRECTORY / "sample.docx"))


def write_xlsx() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Batches"
    worksheet.append(["batch", "status", "yield"])
    worksheet.append(["B-001", "released", 98.4])
    worksheet.append(["B-002", "review", 96.1])
    output = BytesIO()
    workbook.save(output)
    (FIXTURE_DIRECTORY / "sample.xlsx").write_bytes(output.getvalue())


if __name__ == "__main__":
    (FIXTURE_DIRECTORY / "sample.pdf").write_bytes(build_pdf())
    write_docx()
    write_xlsx()
