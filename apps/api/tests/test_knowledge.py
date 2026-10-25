from __future__ import annotations

import asyncio
import json
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from docx import Document as OpenDocument
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.config import Settings
from app.db.base import Base
from app.db.models import IngestionJob, KnowledgeCompilationJob, SourceRecord, WikiChange
from app.ingestion.service import IngestionService
from app.knowledge.compiler import SYSTEM_PROMPT, KnowledgeCompiler
from app.knowledge.errors import InvalidKnowledgeProposalError, WikiConflictError
from app.knowledge.store import WikiStore
from app.knowledge.types import (
    ClaimClassification,
    ClaimStatus,
    CompactKnowledgeProposal,
    KnowledgeProposal,
    PageKind,
    ProvenanceRef,
    WikiClaim,
    WikiPage,
    WikiRelationship,
)
from app.llm import (
    ChatChunk,
    ChatMessage,
    ChatRequest,
    ChatResponse,
    ChatUsage,
    LLMCapabilities,
    LLMHealth,
)
from app.llm.errors import LLMResponseError, LLMTimeoutError
from app.main import create_app

FIXTURES = Path(__file__).parent / "fixtures" / "knowledge"


class ProposalProvider:
    def __init__(self, *outputs: dict[str, object] | str) -> None:
        self.outputs = [
            output if isinstance(output, str) else json.dumps(output) for output in outputs
        ]
        self.requests: list[ChatRequest] = []

    @property
    def capabilities(self) -> LLMCapabilities:
        return LLMCapabilities(
            streaming=True,
            structured_output=True,
            cancellation=True,
            max_context_tokens=4096,
        )

    async def health(self) -> LLMHealth:
        return LLMHealth(healthy=True, provider="fake", model="fake-knowledge")

    async def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        if not self.outputs:
            raise AssertionError("The compiler made an unexpected model call.")
        return ChatResponse(
            message=ChatMessage(role="assistant", content=self.outputs.pop(0)),
            model="fake-knowledge",
            finish_reason="stop",
            usage=ChatUsage(),
        )

    async def stream(self, request: ChatRequest) -> AsyncIterator[ChatChunk]:
        del request
        if False:
            yield ChatChunk(model="fake-knowledge")


class TruncatingProposalProvider(ProposalProvider):
    def __init__(self, *outputs: dict[str, object] | str) -> None:
        super().__init__(*outputs)
        self.simulated_durations: list[float] = []

    @property
    def capabilities(self) -> LLMCapabilities:
        return LLMCapabilities(
            streaming=True,
            structured_output=True,
            cancellation=True,
            max_context_tokens=8192,
        )

    async def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        if not self.outputs:
            raise AssertionError("The compiler made an unexpected model call.")
        # Deterministically model the measured Qwen floor: 8.7 seconds of prompt work plus
        # generation at 18.6 tokens/second. The former 180-second shared deadline would reject the
        # 4,096-token retry; the compilation-specific deadline must contain it without sleeping.
        simulated_duration = 8.7 + request.max_tokens / 18.6
        self.simulated_durations.append(simulated_duration)
        if request.timeout_seconds is not None and simulated_duration >= request.timeout_seconds:
            raise LLMTimeoutError("simulated bounded request timeout")
        truncated = len(self.requests) == 1
        completion_tokens = request.max_tokens if truncated else 400
        return ChatResponse(
            message=ChatMessage(role="assistant", content=self.outputs.pop(0)),
            model="fake-knowledge",
            finish_reason="length" if truncated else "stop",
            usage=ChatUsage(
                prompt_tokens=512,
                completion_tokens=completion_tokens,
                total_tokens=512 + completion_tokens,
            ),
        )


class TimeoutThenProposalProvider(ProposalProvider):
    async def chat(self, request: ChatRequest) -> ChatResponse:
        if not self.requests:
            self.requests.append(request)
            raise LLMTimeoutError("simulated bounded request timeout")
        return await super().chat(request)


class ResponseErrorOnceProvider(ProposalProvider):
    async def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        if len(self.requests) == 1:
            raise LLMResponseError("simulated structured response contract failure")
        if not self.outputs:
            raise AssertionError("The compiler made an unexpected model call.")
        return ChatResponse(
            message=ChatMessage(role="assistant", content=self.outputs.pop(0)),
            model="fake-knowledge",
            finish_reason="stop",
            usage=ChatUsage(),
        )


class ReducedContextProposalProvider(ProposalProvider):
    @property
    def capabilities(self) -> LLMCapabilities:
        return LLMCapabilities(
            streaming=True,
            structured_output=True,
            cancellation=True,
            max_context_tokens=3000,
        )


class AlwaysTruncatingProposalProvider(ProposalProvider):
    @property
    def capabilities(self) -> LLMCapabilities:
        return LLMCapabilities(
            streaming=True,
            structured_output=True,
            cancellation=True,
            max_context_tokens=8192,
        )

    async def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        return ChatResponse(
            message=ChatMessage(role="assistant", content='{"pages":['),
            model="fake-knowledge",
            finish_reason="length",
            usage=ChatUsage(
                prompt_tokens=512,
                completion_tokens=request.max_tokens,
                total_tokens=512 + request.max_tokens,
            ),
        )


class BlockingProposalProvider(ProposalProvider):
    def __init__(self) -> None:
        super().__init__()
        self.started = asyncio.Event()

    async def chat(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        self.started.set()
        await asyncio.Event().wait()
        raise AssertionError("The cancelled compilation unexpectedly resumed.")


def build_app(tmp_path: Path, provider: ProposalProvider) -> FastAPI:
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    settings = Settings(
        _env_file=None,
        database_url="sqlite+pysqlite://",
        llm_model="fake-knowledge",
        authentication_required=False,
        source_store_path=tmp_path / "sources",
        knowledge_path=tmp_path / "knowledge",
        knowledge_compilation_batch_chars=2_000,
    )
    return create_app(settings=settings, provider=provider, engine=engine)


async def upload_and_compile(
    client: AsyncClient,
    filename: str,
    content: bytes,
    content_type: str = "text/markdown",
) -> tuple[dict[str, Any], dict[str, Any]]:
    uploaded = await client.post(
        "/api/v1/documents",
        files={"file": (filename, content, content_type)},
    )
    assert uploaded.status_code == 201, uploaded.text
    compiled = await client.post(
        f"/api/v1/document-versions/{uploaded.json()['document_version_id']}/compile"
    )
    assert compiled.status_code == 200, compiled.text
    return uploaded.json(), compiled.json()


def acme_proposal(*, headquarters: str = "Valletta") -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "entity",
                "title": "Acme Corporation",
                "aliases": ["ACME"],
                "summary": "A medicine distribution company.",
                "claims": [
                    {
                        "key": "headquarters",
                        "text": f"Acme Corporation is headquartered in {headquarters}.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [
                    {
                        "relation": "uses",
                        "target_title": "Cold Chain",
                        "source_record_ordinals": [1],
                    }
                ],
            },
            {
                "type": "concept",
                "title": "Cold Chain",
                "aliases": ["Temperature-controlled distribution"],
                "summary": "Controls for temperature-sensitive distribution.",
                "claims": [
                    {
                        "key": "purpose",
                        "text": "Cold-chain controls protect temperature-sensitive medicine.",
                        "classification": "synthesis",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [
                    {
                        "relation": "used_by",
                        "target_title": "Acme Corporation",
                        "source_record_ordinals": [1],
                    }
                ],
            },
        ],
        "warnings": [],
    }


def standard_proposal() -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "entity",
                "title": "ACME",
                "aliases": ["Acme Corporation"],
                "summary": "A medicine distribution company.",
                "claims": [
                    {
                        "key": "quality_review",
                        "text": "ACME completed a Good Distribution Practice review.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [
                    {
                        "relation": "uses",
                        "target_title": "Cold Chain",
                        "source_record_ordinals": [1],
                    }
                ],
            },
            {
                "type": "topic",
                "title": "Good Distribution Practice",
                "aliases": ["GDP"],
                "summary": "Quality requirements for medicine distribution.",
                "claims": [
                    {
                        "key": "scope",
                        "text": "GDP review covers temperature monitoring and excursions.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [
                    {
                        "relation": "applies_to",
                        "target_title": "ACME",
                        "source_record_ordinals": [1],
                    }
                ],
            },
        ],
        "warnings": [],
    }


def conflict_proposal() -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "entity",
                "title": "ACME",
                "aliases": ["Acme Corporation"],
                "summary": "A medicine distribution company.",
                "claims": [
                    {
                        "key": "headquarters",
                        "text": "Acme Corporation is headquartered in Milan.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [],
            }
        ],
        "warnings": ["Source conflicts with an existing headquarters claim."],
    }


def project_aurora_proposal() -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "topic",
                "title": "Project Aurora",
                "aliases": [],
                "topics": [],
                "summary": "A governed release project.",
                "claims": [
                    {
                        "key": "release_window",
                        "text": "Project Aurora releases on Friday at 14:00 UTC.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [],
            }
        ],
        "warnings": [],
    }


def contaminated_skills_proposal() -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "topic",
                "title": "Project Aurora",
                "aliases": [],
                "topics": [],
                "summary": "A governed release project.",
                "claims": [
                    {
                        "key": "launch_review",
                        "text": "Project Aurora assigns its launch review to Mira Chen.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [],
            }
        ],
        "warnings": [],
    }


def federico_skills_proposal() -> dict[str, object]:
    pages: list[dict[str, object]] = [
        {
            "type": "entity",
            "title": "Avery Morgan",
            "aliases": ["Federico"],
            "topics": ["software development"],
            "summary": "A designer and fullstack developer.",
            "claims": [
                {
                    "key": "technical_context",
                    "text": "Avery Morgan works with TypeScript, JavaScript, and SQL.",
                    "classification": "source_fact",
                    "source_record_ordinals": [1],
                }
            ],
            "relationships": [
                {
                    "relation": "uses",
                    "target_title": title,
                    "source_record_ordinals": [1],
                }
                for title in ("TypeScript", "JavaScript", "SQL")
            ],
        }
    ]
    pages.extend(
        {
            "type": "concept",
            "title": title,
            "aliases": [],
            "topics": ["programming"],
            "summary": f"{title} in Avery Morgan's technical context.",
            "claims": [
                {
                    "key": "used_by",
                    "text": f"Avery Morgan uses {title} in their work context.",
                    "classification": "source_fact",
                    "source_record_ordinals": [1],
                }
            ],
            "relationships": [],
        }
        for title in ("TypeScript", "JavaScript", "SQL")
    )
    return {"pages": pages, "warnings": []}


def named_subjects_proposal(
    *subjects: tuple[str, str, tuple[str, ...]],
    relationships: dict[str, list[dict[str, object]]] | None = None,
) -> dict[str, object]:
    return {
        "pages": [
            {
                "type": kind,
                "title": title,
                "aliases": list(aliases),
                "topics": [],
                "summary": f"Source-backed information about {title}.",
                "claims": [
                    {
                        "key": "source_description",
                        "text": f"The source explicitly discusses {title}.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": (relationships or {}).get(title, []),
            }
            for title, kind, aliases in subjects
        ],
        "warnings": [],
    }


def fabrizio_relationship_proposal(
    colleague_text: str = "Fabrizio was a colleague of Federico La Terra.",
) -> dict[str, object]:
    return {
        "pages": [
            {
                "type": "entity",
                "title": "Fabrizio Bonaccorsi",
                "aliases": [],
                "topics": [],
                "summary": "Fabrizio Bonaccorsi is an Italian web developer.",
                "claims": [
                    {
                        "key": "occupation",
                        "text": "Fabrizio Bonaccorsi is an Italian web developer.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    },
                    {
                        "key": "employment",
                        "text": "Fabrizio Bonaccorsi works at ABC Company.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    },
                    {
                        "key": "professional_relationship",
                        "text": colleague_text,
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    },
                ],
                # Qwen may retain both explicit facts as claims without proposing relationships.
                "relationships": [],
            },
            {
                "type": "entity",
                "title": "ABC Company",
                "aliases": [],
                "topics": [],
                "summary": "ABC Company employs Fabrizio Bonaccorsi.",
                "claims": [
                    {
                        "key": "employee",
                        "text": "Fabrizio Bonaccorsi works at ABC Company.",
                        "classification": "source_fact",
                        "source_record_ordinals": [1],
                    }
                ],
                "relationships": [],
            },
        ],
        "warnings": [],
    }


def compact_fabrizio_relationship_proposal(
    colleague_text: str = "Fabrizio was a colleague of Federico La Terra.",
) -> dict[str, object]:
    return {
        "v": 2,
        "p": [
            {"k": "entity", "t": "Fabrizio Bonaccorsi"},
            {"k": "entity", "t": "ABC Company"},
        ],
        "f": [
            {
                "on": [{"p": 0, "k": "occupation"}],
                "x": "Fabrizio Bonaccorsi is an Italian web developer.",
                "s": [1],
            },
            {
                "on": [
                    {"p": 0, "k": "employment"},
                    {"p": 1, "k": "employee"},
                ],
                "x": "Fabrizio Bonaccorsi works at ABC Company.",
                "s": [1],
            },
            {
                "on": [{"p": 0, "k": "professional_relationship"}],
                "x": colleague_text,
                "s": [1],
            },
        ],
    }


def equivalent_people_proposal(
    *, introduction_ordinal: int, employment_ordinal: int, colleague_ordinal: int
) -> dict[str, object]:
    proposal = fabrizio_relationship_proposal()
    pages = proposal["pages"]
    assert isinstance(pages, list)
    fabrizio = pages[0]
    abc = pages[1]
    assert isinstance(fabrizio, dict) and isinstance(abc, dict)
    claims = fabrizio["claims"]
    assert isinstance(claims, list)
    claims[0]["source_record_ordinals"] = [introduction_ordinal]
    claims[1]["source_record_ordinals"] = [employment_ordinal]
    claims[2]["source_record_ordinals"] = [colleague_ordinal]
    abc_claims = abc["claims"]
    assert isinstance(abc_claims, list)
    abc_claims[0]["source_record_ordinals"] = [employment_ordinal]
    return proposal


def compact_equivalent_people_proposal(
    *, introduction_ordinal: int, employment_ordinal: int, colleague_ordinal: int
) -> dict[str, object]:
    proposal = compact_fabrizio_relationship_proposal()
    facts = proposal["f"]
    assert isinstance(facts, list)
    for fact, ordinal in zip(
        facts,
        (introduction_ordinal, employment_ordinal, colleague_ordinal),
        strict=True,
    ):
        assert isinstance(fact, dict)
        fact["s"] = [ordinal]
    return proposal


def build_text_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    commands = ["BT /F1 11 Tf 72 740 Td"]
    for index, line in enumerate(escaped.splitlines()):
        if index:
            commands.append("0 -18 Td")
        commands.append(f"({line}) Tj")
    commands.append("ET")
    content = ("\n".join(commands) + "\n").encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"endstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(f"{number} 0 obj\n".encode())
        result.extend(body)
        result.extend(b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(result)


def build_people_fixture(
    format_name: str, *, headed: bool
) -> tuple[bytes, str, str, tuple[int, ...]]:
    paragraphs = (
        "Fabrizio Bonaccorsi is an Italian web developer.",
        "He works at ABC Company.",
        "Fabrizio was a colleague of Federico La Terra.",
        "Fabrizio might be a friend of Federico La Terra and may prefer sushi.",
        'A note quotes "Fabrizio is the owner of ABC Company."',
        "If Fabrizio studied at Malta University, that would be recorded separately.",
        "The unrelated garden paragraph discusses tomato seedlings.",
    )
    lines = (("People",) if headed else ()) + paragraphs
    text = "\n".join(lines)
    if format_name == "txt":
        return text.encode(), "people.txt", "text/plain", (1, 1, 1)
    if format_name == "pdf":
        return build_text_pdf(text), "people.pdf", "application/pdf", (1, 1, 1)
    document = OpenDocument()
    if headed:
        document.add_heading("People", level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    offset = 1 if headed else 0
    return (
        output.getvalue(),
        "people.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        (1 + offset, 2 + offset, 3 + offset),
    )


def test_compact_v2_proposal_expands_without_losing_claim_semantics() -> None:
    legacy = KnowledgeProposal.model_validate(fabrizio_relationship_proposal())
    compact_payload = compact_fabrizio_relationship_proposal()
    compact = CompactKnowledgeProposal.model_validate(compact_payload)
    expanded = compact.expand()

    def claim_semantics(proposal: KnowledgeProposal) -> set[tuple[str, str, str, tuple[int, ...]]]:
        return {
            (page.title, claim.key, claim.text, claim.source_record_ordinals)
            for page in proposal.pages
            for claim in page.claims
        }

    assert compact.contract_version == 2
    assert claim_semantics(expanded) == claim_semantics(legacy)
    assert all(page.summary == page.claims[0].text for page in expanded.pages)
    legacy_size = len(json.dumps(fabrizio_relationship_proposal(), separators=(",", ":")))
    compact_size = len(json.dumps(compact_payload, separators=(",", ":")))
    assert compact_size <= legacy_size * 0.65


def test_compact_v2_proposal_rejects_unknown_page_references() -> None:
    payload = compact_fabrizio_relationship_proposal()
    facts = payload["f"]
    assert isinstance(facts, list)
    fact = facts[0]
    assert isinstance(fact, dict)
    fact["on"] = [{"p": 99, "k": "occupation"}]

    with pytest.raises(ValueError, match="unknown page index"):
        CompactKnowledgeProposal.model_validate(payload)


def test_compact_v2_relationship_expands_to_legacy_relationship() -> None:
    payload = compact_fabrizio_relationship_proposal()
    payload["r"] = [
        {
            "f": 0,
            "r": "works_at",
            "t": "ABC Company",
            "s": [1],
        }
    ]

    expanded = CompactKnowledgeProposal.model_validate(payload).expand()

    assert len(expanded.pages[0].relationships) == 1
    assert expanded.pages[0].relationships[0].model_dump(mode="json") == {
        "relation": "works_at",
        "target_title": "ABC Company",
        "source_record_ordinals": [1],
    }


def test_related_sources_build_one_evolving_wiki_with_conflict_provenance(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(acme_proposal(), standard_proposal(), conflict_proposal())
        app = build_app(tmp_path, provider)
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            first_upload, first = await upload_and_compile(
                client,
                "acme-profile.md",
                (FIXTURES / "acme-profile.md").read_bytes(),
            )
            _, second = await upload_and_compile(
                client,
                "cold-chain-standard.md",
                (FIXTURES / "cold-chain-standard.md").read_bytes(),
            )
            _, third = await upload_and_compile(
                client,
                "acme-audit.md",
                (FIXTURES / "acme-audit.md").read_bytes(),
            )
            pages_response = await client.get("/api/v1/wiki/pages")
            page_index_response = await client.get(
                "/api/v1/wiki/page-index?page=1&page_limit=2&q=acme&kind=entity"
            )
            acme_response = await client.get("/api/v1/wiki/pages/ACME")
            backlink_response = await client.get("/api/v1/wiki/pages/ACME/backlinks")
            graph_response = await client.get("/api/v1/wiki/graph")
            local_graph_response = await client.get("/api/v1/wiki/graph?root=ACME&depth=1")
            lint_response = await client.get("/api/v1/wiki/lint")
            repeated = await client.post(
                f"/api/v1/document-versions/{first_upload['document_version_id']}/compile"
            )
            observed = await client.get(f"/api/v1/knowledge-compilation-jobs/{first['id']}")

        assert first["status"] == "COMPLETED"
        assert len(first["report"]["pages_created"]) == 2
        assert second["status"] == "COMPLETED"
        assert len(second["report"]["pages_created"]) == 1
        assert len(second["report"]["pages_updated"]) == 1
        assert third["status"] == "COMPLETED"
        assert len(third["report"]["conflicts"]) == 1
        assert repeated.json()["id"] == first["id"]
        assert observed.json()["status"] == "COMPLETED"
        assert len(provider.requests) == 3

        pages = pages_response.json()
        assert len(pages) == 3
        assert {page["title"] for page in pages} == {
            "Acme Corporation",
            "Cold Chain",
            "Good Distribution Practice",
        }
        page_index = page_index_response.json()
        assert page_index["page"] == 1
        assert page_index["page_limit"] == 2
        assert page_index["total"] == 1
        assert page_index["total_pages"] == 1
        assert [page["title"] for page in page_index["items"]] == ["Acme Corporation"]
        assert acme_response.json()["title"] == "Acme Corporation"
        assert {page["title"] for page in backlink_response.json()} == {
            "Cold Chain",
            "Good Distribution Practice",
        }
        graph = graph_response.json()
        assert graph["mode"] == "global"
        assert graph["total_nodes"] == 3
        assert graph["total_edges"] == 3
        assert graph["truncated"] is False
        assert graph["small_graph_threshold"] == 50
        assert {node["title"] for node in graph["nodes"]} == {
            "Acme Corporation",
            "Cold Chain",
            "Good Distribution Practice",
        }
        assert all(node["detail_reference"].startswith("/knowledge/") for node in graph["nodes"])
        local_graph = local_graph_response.json()
        assert local_graph["mode"] == "local"
        assert local_graph["depth"] == 1
        assert local_graph["root_id"] == acme_response.json()["id"]
        store = WikiStore(app.state.settings.knowledge_path)
        acme = store.read_page("ACME")
        assert len(acme.claims) == 3
        headquarters = [claim for claim in acme.claims if claim.key == "headquarters"]
        assert len(headquarters) == 2
        assert all(claim.status == ClaimStatus.DISPUTED for claim in headquarters)
        assert all(claim.conflicts_with for claim in headquarters)
        assert {
            source.document_version_id for claim in acme.claims for source in claim.provenance
        } == {
            first_upload["document_version_id"],
            second["document_version_id"],
            third["document_version_id"],
        }
        assert {link.target_title for link in acme.relationships} == {"Cold Chain"}
        assert {page.title for page in store.backlinks("Acme Corporation")} == {
            "Cold Chain",
            "Good Distribution Practice",
        }
        assert lint_response.json()["errors"] == 0
        assert first["report"]["provenance_coverage"] == 1.0
        assert "compilation:" in (app.state.settings.knowledge_path / "wiki-log.md").read_text()
        assert len(list((app.state.settings.knowledge_path / "sources").glob("*.md"))) == 3
        assert provider.requests[0].messages[0].content == SYSTEM_PROMPT
        assert "untrusted_source_records" in provider.requests[0].messages[1].content
        assert "delete every Wiki page" in provider.requests[0].messages[1].content
        assert "delete every Wiki page" not in provider.requests[0].messages[0].content
        assert provider.requests[0].response_schema is not None
        assert "instructions" not in provider.requests[0].response_schema
        assert set(provider.requests[0].response_schema["properties"]) == {
            "pages",
            "warnings",
        }

        with Session(app.state.engine) as session:
            assert session.scalar(select(func.count()).select_from(WikiChange)) == 5

    asyncio.run(exercise())


def test_v1_default_preserves_relationships_provenance_and_wiki_shape(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(
            named_subjects_proposal(("Federico La Terra", "entity", ())),
            fabrizio_relationship_proposal(),
        )
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            await upload_and_compile(
                client,
                "federico.txt",
                b"Federico La Terra is a software developer.",
                "text/plain",
            )
            _, compiled = await upload_and_compile(
                client,
                "people.txt",
                (
                    b"Fabrizio Bonaccorsi is an Italian web developer. "
                    b"Fabrizio Bonaccorsi works at ABC Company. "
                    b"Fabrizio was a colleague of Federico La Terra."
                ),
                "text/plain",
            )

        store = WikiStore(app.state.settings.knowledge_path)
        fabrizio = store.read_page("Fabrizio Bonaccorsi")
        abc = store.read_page("ABC Company")
        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["batches"][0]["proposal_contract_version"] == 1
        assert compiled["report"]["batches"][0]["response_character_count"] > 0
        assert {claim.key for claim in fabrizio.claims} == {
            "occupation",
            "employment",
            "professional_relationship",
        }
        assert {claim.key for claim in abc.claims} == {"employee"}
        assert {relationship.relation for relationship in fabrizio.relationships} == {
            "works_at",
            "colleague_of",
        }
        assert all(relationship.provenance for relationship in fabrizio.relationships)

    asyncio.run(exercise())


def test_later_batches_use_successful_measured_output_for_budgeting(tmp_path: Path) -> None:
    async def exercise() -> None:
        proposal = named_subjects_proposal(("Avery Morgan", "entity", ()))
        provider = ProposalProvider(proposal, proposal)
        app = build_app(tmp_path, provider)
        source = ("Avery Morgan leads operations. " * 80).encode()
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            _, compiled = await upload_and_compile(client, "multi.txt", source, "text/plain")

        assert compiled["status"] == "COMPLETED"
        assert len(provider.requests) == 2
        assert provider.requests[0].max_tokens == 2048
        assert provider.requests[1].max_tokens < provider.requests[0].max_tokens
        attempts = [batch["attempts"][0] for batch in compiled["report"]["batches"]]
        assert [attempt["budget_basis"] for attempt in attempts] == [
            "configured_bootstrap",
            "measured_task_requirements",
        ]
        assert all(len(batch["attempts"]) == 1 for batch in compiled["report"]["batches"])

    asyncio.run(exercise())


def test_bootstrap_output_budget_respects_effective_provider_context(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = ReducedContextProposalProvider(fabrizio_relationship_proposal())
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            _, compiled = await upload_and_compile(
                client,
                "people.txt",
                (
                    b"Fabrizio Bonaccorsi is an Italian web developer. "
                    b"Fabrizio Bonaccorsi works at ABC Company. "
                    b"Fabrizio was a colleague of Federico La Terra."
                ),
                "text/plain",
            )

        assert compiled["status"] == "COMPLETED"
        assert provider.requests[0].max_tokens == 1500
        attempt = compiled["report"]["batches"][0]["attempts"][0]
        assert attempt["budget_basis"] == "effective_context_bootstrap"

    asyncio.run(exercise())


def test_compilation_blocks_unrelated_existing_pages_and_builds_explicit_skill_links(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(
            project_aurora_proposal(),
            contaminated_skills_proposal(),
            federico_skills_proposal(),
        )
        app = build_app(tmp_path, provider)
        transport = ASGITransport(app=app)
        skills = (
            b"Avery Morgan uses TypeScript for frontend and backend development.\n"
            b"Avery Morgan works in the JavaScript ecosystem.\n"
            b"Avery Morgan's work context includes relational databases based on SQL.\n"
        )
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            _, first = await upload_and_compile(
                client,
                "project-aurora.md",
                b"Project Aurora releases on Friday at 14:00 UTC.",
            )
            wiki_before_rejection = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }
            uploaded = await client.post(
                "/api/v1/documents",
                files={
                    "file": (
                        "Federico_La_Terra_Skills.txt",
                        skills,
                        "text/plain",
                    )
                },
            )
            assert uploaded.status_code == 201, uploaded.text
            version_id = uploaded.json()["document_version_id"]
            rejected = await client.post(f"/api/v1/document-versions/{version_id}/compile")
            aurora_after_rejection = await client.get("/api/v1/wiki/pages/Project%20Aurora")
            wiki_after_rejection = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }
            retried = await client.post(f"/api/v1/document-versions/{version_id}/compile")
            local_graph = await client.get(
                "/api/v1/wiki/graph",
                params={"root": "Avery Morgan", "depth": 1},
            )

        assert first["status"] == "COMPLETED"
        assert rejected.status_code == 200
        rejected_job = rejected.json()
        assert rejected_job["status"] == "FAILED"
        assert rejected_job["error"] == (
            "Knowledge proposal attempted to update existing pages not named in source "
            "records: Project Aurora."
        )
        assert aurora_after_rejection.status_code == 200
        assert {claim["text"] for claim in aurora_after_rejection.json()["claims"]} == {
            "Project Aurora releases on Friday at 14:00 UTC."
        }
        assert wiki_after_rejection == wiki_before_rejection

        request_payload = json.loads(
            provider.requests[1].messages[1].content.removeprefix("UNTRUSTED_SOURCE_JSON\n")
        )
        assert "existing_wiki_pages" not in request_payload
        assert request_payload["existing_page_identities"] == []

        retried_job = retried.json()
        assert retried_job["status"] == "COMPLETED"
        assert retried_job["attempt_count"] == 2
        assert retried_job["report"]["links_created"] == 3
        graph = local_graph.json()
        assert graph["mode"] == "local"
        assert {node["title"] for node in graph["nodes"]} == {
            "Avery Morgan",
            "JavaScript",
            "SQL",
            "TypeScript",
        }
        assert len(graph["edges"]) == 3
        assert all("uses" in edge["relations"] for edge in graph["edges"])

    asyncio.run(exercise())


def test_reprocess_preserves_source_record_identity_and_forces_a_new_compilation_attempt(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(acme_proposal(), acme_proposal())
        app = build_app(tmp_path, provider)
        transport = ASGITransport(app=app)
        source = b"Acme Corporation is headquartered in Valletta and uses Cold Chain."
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            uploaded = await client.post(
                "/api/v1/documents",
                files={"file": ("acme.md", source, "text/markdown")},
            )
            assert uploaded.status_code == 201, uploaded.text
            body = uploaded.json()
            version_id = uuid.UUID(body["document_version_id"])
            first = await client.post(f"/api/v1/document-versions/{version_id}/compile")
            assert first.status_code == 200, first.text
            assert first.json()["status"] == "COMPLETED"
            wiki_before_reprocess = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }

            with Session(app.state.engine) as session:
                before_ids = tuple(
                    session.scalars(
                        select(SourceRecord.id)
                        .where(SourceRecord.document_version_id == version_id)
                        .order_by(SourceRecord.ordinal)
                    )
                )

            reprocessed = await client.post(
                f"/api/v1/ingestion-jobs/{body['ingestion_job_id']}/reprocess"
            )
            original = await client.get(f"/api/v1/document-versions/{version_id}/original")
            second = await client.post(f"/api/v1/document-versions/{version_id}/compile")

        assert reprocessed.status_code == 200, reprocessed.text
        assert reprocessed.json()["status"] == "READY_FOR_COMPILATION"
        assert original.content == source
        assert second.status_code == 200, second.text
        assert second.json()["status"] == "COMPLETED"
        assert second.json()["attempt_count"] == 2
        assert len(provider.requests) == 2
        with Session(app.state.engine) as session:
            after_ids = tuple(
                session.scalars(
                    select(SourceRecord.id)
                    .where(SourceRecord.document_version_id == version_id)
                    .order_by(SourceRecord.ordinal)
                )
            )
            compilation = session.scalar(
                select(KnowledgeCompilationJob).where(
                    KnowledgeCompilationJob.document_version_id == version_id
                )
            )
            assert compilation is not None
            assert compilation.status == "COMPLETED"
        assert after_ids == before_ids
        wiki_after_reprocess = {
            path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
            for path in app.state.settings.knowledge_path.rglob("*")
            if path.is_file()
        }
        assert wiki_after_reprocess == wiki_before_reprocess

    asyncio.run(exercise())


class FailOnceStore(WikiStore):
    def __init__(self, root: Path) -> None:
        super().__init__(root)
        self.fail_next_page_write = True

    def _write_page(self, path: Path, page: WikiPage) -> None:
        super()._write_page(path, page)
        if self.fail_next_page_write:
            self.fail_next_page_write = False
            raise OSError("simulated interruption")


class ConflictOnceStore(WikiStore):
    def __init__(self, root: Path) -> None:
        super().__init__(root)
        self.fail_next_page_write = True

    def _write_page(self, path: Path, page: WikiPage) -> None:
        super()._write_page(path, page)
        if self.fail_next_page_write:
            self.fail_next_page_write = False
            raise WikiConflictError("simulated stale proposal conflict")


def test_interrupted_compilation_rolls_back_and_reuses_validated_proposal(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(fabrizio_relationship_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        store = FailOnceStore(settings.knowledge_path)
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="people.md",
                declared_content_type="text/markdown",
                data=(
                    b"Fabrizio Bonaccorsi is an Italian web developer. "
                    b"Fabrizio Bonaccorsi works at ABC Company. "
                    b"Fabrizio was a colleague of Federico La Terra."
                ),
            )
            compiler = KnowledgeCompiler(
                session=session,
                settings=settings,
                provider=provider,
                store=store,
            )
            failed = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert failed.status == "FAILED"
            assert store.list_pages() == ()
            assert failed.proposal is not None
            assert "pages" in failed.proposal
            assert "v" not in failed.proposal
            recovered = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert recovered.status == "COMPLETED"
            assert recovered.attempt_count == 2
            assert len(store.list_pages()) == 2
            assert len(provider.requests) == 1

    asyncio.run(exercise())


def test_invalid_structured_model_output_is_failed_and_retryable(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = ProposalProvider("not valid JSON", acme_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiler = KnowledgeCompiler(session=session, settings=settings, provider=provider)
            failed = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert failed.status == "FAILED"
            assert failed.proposal is None
            assert failed.error == (
                "Local model output did not match the structured knowledge proposal schema."
            )
            assert WikiStore(settings.knowledge_path).list_pages() == ()
            assert len(provider.requests) == 1
            retried = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert retried.status == "COMPLETED"
            assert retried.attempt_count == 2
            assert len(provider.requests) == 2
            assert all(
                set(request.response_schema["properties"]) == {"pages", "warnings"}
                for request in provider.requests
                if request.response_schema is not None
            )

    asyncio.run(exercise())


def test_provider_response_error_fails_without_automatic_fallback(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = ResponseErrorOnceProvider(acme_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiler = KnowledgeCompiler(session=session, settings=settings, provider=provider)
            failed = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert failed.status == "FAILED"
            assert failed.error == (
                "Knowledge compilation failed because the local model returned an invalid response."
            )
            assert failed.proposal is None
            assert WikiStore(settings.knowledge_path).list_pages() == ()
            assert len(provider.requests) == 1
            compiled = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            compiled_status = compiled.status
            compiled_report = compiled.report

        assert compiled_status == "COMPLETED"
        assert len(provider.requests) == 2
        assert all(
            request.response_schema is not None
            and set(request.response_schema["properties"]) == {"pages", "warnings"}
            for request in provider.requests
        )
        assert compiled_report is not None
        batch = compiled_report["batches"][0]
        assert batch["proposal_contract_version"] == 1
        assert len(batch["attempts"]) == 1
        assert batch["attempts"][0]["retry_cause"] is None

    asyncio.run(exercise())


def test_compact_v2_model_response_is_rejected_without_fallback(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(compact_fabrizio_relationship_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="people.md",
                declared_content_type="text/markdown",
                data=(
                    b"Fabrizio Bonaccorsi is an Italian web developer. "
                    b"Fabrizio Bonaccorsi works at ABC Company. "
                    b"Fabrizio was a colleague of Federico La Terra."
                ),
            )
            failed = await KnowledgeCompiler(
                session=session, settings=settings, provider=provider
            ).compile(owner_id=owner.id, document_version_id=source.version.id)
            failed_status = failed.status
            failed_error = failed.error

        assert failed_status == "FAILED"
        assert failed_error == (
            "Local model output did not match the structured knowledge proposal schema."
        )
        assert len(provider.requests) == 1
        assert WikiStore(settings.knowledge_path).list_pages() == ()

    asyncio.run(exercise())


def test_valid_short_fixture_uses_one_v1_request_without_fallback(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(acme_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiled = await KnowledgeCompiler(
                session=session, settings=settings, provider=provider
            ).compile(owner_id=owner.id, document_version_id=source.version.id)
            status = compiled.status
            report = compiled.report

        assert status == "COMPLETED"
        assert len(provider.requests) == 1
        request = provider.requests[0]
        assert request.messages[0].content == SYSTEM_PROMPT
        assert request.max_tokens == 2048
        assert request.response_schema is not None
        assert set(request.response_schema["properties"]) == {"pages", "warnings"}
        assert report is not None
        batch = report["batches"][0]
        assert batch["proposal_contract_version"] == 1
        assert len(batch["attempts"]) == 1
        assert batch["attempts"][0]["retry_cause"] is None

    asyncio.run(exercise())


def test_wiki_conflict_discards_cached_proposal_before_user_retry(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(
            fabrizio_relationship_proposal(),
            fabrizio_relationship_proposal(),
        )
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        store = ConflictOnceStore(settings.knowledge_path)
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="people.md",
                declared_content_type="text/markdown",
                data=(
                    b"Fabrizio Bonaccorsi is an Italian web developer. "
                    b"Fabrizio Bonaccorsi works at ABC Company. "
                    b"Fabrizio was a colleague of Federico La Terra."
                ),
            )
            compiler = KnowledgeCompiler(
                session=session,
                settings=settings,
                provider=provider,
                store=store,
            )
            failed = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert failed.status == "FAILED"
            assert failed.error == "Knowledge compilation failed: WikiConflictError."
            assert failed.proposal is None
            assert failed.last_completed_stage is None
            assert store.list_pages() == ()

            recovered = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            assert recovered.status == "COMPLETED"
            assert recovered.attempt_count == 2
            assert len(provider.requests) == 2

    asyncio.run(exercise())


def test_truncated_structured_output_retries_with_larger_budget(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = TruncatingProposalProvider('{"pages":[', acme_proposal())
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiler = KnowledgeCompiler(
                session=session,
                settings=settings,
                provider=provider,
            )
            compiled = await compiler.compile(
                owner_id=owner.id, document_version_id=source.version.id
            )
            status = compiled.status

        assert status == "COMPLETED"
        assert [request.max_tokens for request in provider.requests] == [2048, 4096]
        assert [request.timeout_seconds for request in provider.requests] == [300.0, 300.0]
        assert provider.simulated_durations[1] > 180
        assert provider.simulated_durations[1] < 300

    asyncio.run(exercise())


def test_timed_out_compilation_is_truthful_atomic_and_retryable(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = TimeoutThenProposalProvider(acme_proposal())
        app = build_app(tmp_path, provider)
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            uploaded = await client.post(
                "/api/v1/documents",
                files={
                    "file": (
                        "acme-profile.md",
                        (FIXTURES / "acme-profile.md").read_bytes(),
                        "text/markdown",
                    )
                },
            )
            assert uploaded.status_code == 201
            version_id = uploaded.json()["document_version_id"]

            failed = await client.post(f"/api/v1/document-versions/{version_id}/compile")
            assert failed.status_code == 200
            assert failed.json()["status"] == "FAILED"
            assert failed.json()["attempt_count"] == 1
            assert failed.json()["last_completed_stage"] is None
            assert failed.json()["error"] == (
                "Knowledge compilation timed out while waiting for the local model. "
                "No knowledge changes were applied; retry compilation when the model is ready."
            )
            assert WikiStore(app.state.settings.knowledge_path).list_pages() == ()

            retried = await client.post(f"/api/v1/document-versions/{version_id}/compile")
            assert retried.status_code == 200
            assert retried.json()["status"] == "COMPLETED"
            assert retried.json()["attempt_count"] == 2
            assert len(WikiStore(app.state.settings.knowledge_path).list_pages()) == 2

    asyncio.run(exercise())


def test_second_truncated_structured_response_fails_after_one_retry(tmp_path: Path) -> None:
    async def exercise() -> None:
        provider = AlwaysTruncatingProposalProvider()
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            failed = await KnowledgeCompiler(
                session=session, settings=settings, provider=provider
            ).compile(owner_id=owner.id, document_version_id=source.version.id)

        assert failed.status == "FAILED"
        assert failed.error == (
            "Local model output exceeded the structured knowledge proposal token budget."
        )
        assert [request.max_tokens for request in provider.requests] == [2048, 4096]
        assert WikiStore(settings.knowledge_path).list_pages() == ()

    asyncio.run(exercise())


def test_cancelled_compilation_records_retryable_failure_without_wiki_changes(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        provider = BlockingProposalProvider()
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiler = KnowledgeCompiler(session=session, settings=settings, provider=provider)
            task = asyncio.create_task(
                compiler.compile(owner_id=owner.id, document_version_id=source.version.id)
            )
            await provider.started.wait()
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task

            session.expire_all()
            failed = session.scalar(
                select(KnowledgeCompilationJob).where(
                    KnowledgeCompilationJob.document_version_id == source.version.id
                )
            )
            assert failed is not None
            assert failed.status == "FAILED"
            assert failed.error == (
                "Knowledge compilation was cancelled while waiting for the local model. "
                "No knowledge changes were applied; retry compilation when ready."
            )
            assert failed.last_completed_stage is None
            assert WikiStore(settings.knowledge_path).list_pages() == ()

    asyncio.run(exercise())


def test_single_record_compilation_normalizes_model_ordinals(tmp_path: Path) -> None:
    async def exercise() -> None:
        proposal = json.loads(
            json.dumps(acme_proposal()).replace(
                '"source_record_ordinals": [1]', '"source_record_ordinals": [2]'
            )
        )
        provider = ProposalProvider(proposal)
        app = build_app(tmp_path, provider)
        settings = app.state.settings
        with Session(app.state.engine) as session:
            ingestion = IngestionService(session=session, settings=settings)
            owner = ingestion.get_or_create_local_user()
            source = ingestion.ingest(
                owner=owner,
                filename="acme-profile.md",
                declared_content_type="text/markdown",
                data=(FIXTURES / "acme-profile.md").read_bytes(),
            )
            compiled = await KnowledgeCompiler(
                session=session,
                settings=settings,
                provider=provider,
            ).compile(owner_id=owner.id, document_version_id=source.version.id)
            status = compiled.status
            report = compiled.report

        assert status == "COMPLETED"
        assert report is not None
        assert report["warnings"] == [
            "Model source ordinals were normalized to the only source record in this batch."
        ]

    asyncio.run(exercise())


def fake_provenance() -> ProvenanceRef:
    return ProvenanceRef(
        source_record_id=str(uuid.uuid4()),
        document_id=str(uuid.uuid4()),
        document_version_id=str(uuid.uuid4()),
        original_filename="source.md",
        sha256="a" * 64,
        locator={"heading_path": ["Source"]},
        parser_name="markdown",
        parser_version="1",
        compilation_version="test/1",
        compiled_at=datetime.now(UTC),
    )


def fake_page(title: str, kind: PageKind = PageKind.CONCEPT) -> WikiPage:
    page_id = str(uuid.uuid4())
    provenance = (fake_provenance(),)
    return WikiPage(
        id=page_id,
        type=kind,
        title=title,
        aliases=(),
        summary=f"Summary for {title}.",
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
        compilation_version="test/1",
        provenance=provenance,
        claims=(
            WikiClaim(
                id=str(uuid.uuid4()),
                key="definition",
                text=f"{title} is a test concept.",
                classification=ClaimClassification.SOURCE_FACT,
                provenance=provenance,
            ),
        ),
    )


def test_existing_page_authorization_is_bound_to_each_cited_record() -> None:
    existing = fake_page("Project Aurora")
    proposal = KnowledgeProposal.model_validate(project_aurora_proposal())
    version_id = uuid.uuid4()
    records = tuple(
        (
            SourceRecord(
                id=uuid.uuid4(),
                document_version_id=version_id,
                ordinal=ordinal,
                content=content,
                locator={"line_start": ordinal, "line_end": ordinal},
                parser_name="test",
                parser_version="1",
            ),
            content,
        )
        for ordinal, content in (
            (1, "Unrelated garden schedule."),
            (2, "Project Aurora appears only here."),
        )
    )

    with pytest.raises(InvalidKnowledgeProposalError, match="Project Aurora"):
        KnowledgeCompiler._validate_existing_page_grounding(
            proposal,
            batch=records,
            existing_pages=(existing,),
        )


def test_typed_wiki_operations_and_lint_detect_broken_links_and_orphans(
    tmp_path: Path,
) -> None:
    store = WikiStore(tmp_path / "knowledge")
    alpha = fake_page("Alpha")
    beta = fake_page("Beta")
    store.create_page(alpha)
    store.create_page(beta)
    assert {page.title for page in store.list_pages()} == {"Alpha", "Beta"}
    assert store.read_page(alpha.id).title == "Alpha"

    alpha = alpha.model_copy(update={"aliases": ("First",), "updated_at": datetime.now(UTC)})
    store.update_page(alpha)
    assert store.resolve_page("first").id == alpha.id  # type: ignore[union-attr]
    store.link_pages("Alpha", "Beta", "relates_to")
    assert [page.title for page in store.backlinks("Beta")] == ["Alpha"]
    store.rebuild_index()
    assert store.lint().errors == 0

    orphan = fake_page("Orphan")
    store.create_page(orphan)
    lint = store.lint()
    assert any(issue.code == "orphan-page" and "Orphan" in issue.message for issue in lint.issues)

    broken = WikiRelationship(
        relation="references",
        target_id=str(uuid.uuid4()),
        target_title="Missing Page",
        provenance=alpha.provenance,
    )
    alpha = store.read_page("Alpha").model_copy(update={"relationships": (broken,)})
    store.update_page(alpha)
    lint = store.lint()
    assert {issue.code for issue in lint.issues} >= {
        "broken-wikilink",
        "broken-relationship",
    }


def test_compilation_job_table_is_registered() -> None:
    assert KnowledgeCompilationJob.__tablename__ in Base.metadata.tables


def test_legacy_wiki_metadata_is_copied_without_deleting_original(tmp_path: Path) -> None:
    root = tmp_path / "knowledge"
    legacy = root / ".pharmatrace-wiki"
    legacy.mkdir(parents=True)
    (legacy / "state.json").write_text('{"schema_version": 1, "revision": 7}\n')

    store = WikiStore(root)

    assert store.internal_root == root / ".knemora"
    assert (store.internal_root / "state.json").read_text() == (legacy / "state.json").read_text()
    assert legacy.exists()


@pytest.mark.parametrize(
    ("source", "colleague_text"),
    (
        (
            (
                "FABRIZIO BONACCORSI BIO\n\n"
                "Fabrizio Bonaccorsi is an Italian web developer who specialises in the PHP "
                "programming language and the Magento framework.\n"
                "He works at ABC Company and lives in Cesena with his family.\n\n"
                "PROFESSIONAL RELATIONSHIP\n"
                "Fabrizio was a colleague of Federico La Terra; they are also very good friends "
                "who share a passion for eating sushi.\n"
            ),
            "Fabrizio was a colleague of Federico La Terra.",
        ),
        (
            """Fabrizio Bonaccorsi is an Italian web developer.
He works at ABC Company.
Fabrizio was a colleague of Federico La Terra.
""",
            "Fabrizio was a colleague of Federico La Terra.",
        ),
        (
            (
                "  Fabrizio Bonaccorsi is an Italian web developer.  \n\n\n"
                "He works at ABC Company.   \n\n"
                "Fabrizio was a colleague of Federico La Terra.  \n"
            ),
            "Fabrizio was a colleague of Federico La Terra.",
        ),
        (
            """Fabrizio Bonaccorsi is an Italian web developer.
He works at ABC Company.
Fabrizio is a colleague of Federico La Terra.
""",
            "Fabrizio is a colleague of Federico La Terra.",
        ),
        (
            """Fabrizio Bonaccorsi is an Italian web developer.
He works at ABC Company.
Fabrizio is a former colleague of Federico La Terra.
""",
            "Fabrizio is a former colleague of Federico La Terra.",
        ),
    ),
)
def test_natural_prose_recovers_historical_colleague_and_pronoun_employment_edges(
    tmp_path: Path, source: str, colleague_text: str
) -> None:
    async def exercise() -> None:
        federico = named_subjects_proposal(("Federico La Terra", "entity", ()))
        relationship_proposal = fabrizio_relationship_proposal(colleague_text)
        provider = ProposalProvider(federico, relationship_proposal, relationship_proposal)
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            await upload_and_compile(
                client,
                "federico.md",
                b"Federico La Terra is a software developer.",
            )
            federico_path = next(
                path
                for path in app.state.settings.knowledge_path.rglob("*.md")
                if path.stem == "federico-la-terra"
            )
            federico_before = federico_path.read_bytes()
            uploaded, compiled = await upload_and_compile(
                client, "fabrizio.txt", source.encode(), "text/plain"
            )
            graph = (await client.get("/api/v1/wiki/graph")).json()
            repeated = await client.post(
                f"/api/v1/document-versions/{uploaded['document_version_id']}/compile"
            )
            wiki_before_reprocess = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }
            reprocessed = await client.post(
                f"/api/v1/ingestion-jobs/{uploaded['ingestion_job_id']}/reprocess"
            )
            recompiled = await client.post(
                f"/api/v1/document-versions/{uploaded['document_version_id']}/compile"
            )

        store = WikiStore(app.state.settings.knowledge_path)
        fabrizio = store.read_page("Fabrizio Bonaccorsi")
        federico_page = store.read_page("Federico La Terra")
        abc = store.read_page("ABC Company")
        relations = {
            (page.id, relationship.relation, relationship.target_id)
            for _, page in store.page_snapshot()
            for relationship in page.relationships
        }

        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["links_created"] == 2
        assert relations == {
            (fabrizio.id, "colleague_of", federico_page.id),
            (fabrizio.id, "works_at", abc.id),
        }
        assert (
            next(
                claim.text for claim in fabrizio.claims if claim.key == "professional_relationship"
            )
            == colleague_text
        )
        assert federico_path.read_bytes() == federico_before
        assert all(
            relationship.relation not in {"friend_of", "likes_sushi"}
            for _, page in store.page_snapshot()
            for relationship in page.relationships
        )
        assert graph["total_edges"] == 2
        graph_edges = {
            (edge["source"], edge["target"]): set(edge["relations"]) for edge in graph["edges"]
        }
        assert set(graph_edges) == {
            (fabrizio.id, abc.id),
            (fabrizio.id, federico_page.id),
        }
        assert "works_at" in graph_edges[(fabrizio.id, abc.id)]
        assert "colleague_of" in graph_edges[(fabrizio.id, federico_page.id)]
        assert repeated.json()["id"] == compiled["id"]
        assert reprocessed.status_code == 200
        assert recompiled.json()["report"]["pages_created"] == []
        assert recompiled.json()["report"]["pages_updated"] == []
        wiki_after_reprocess = {
            path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
            for path in app.state.settings.knowledge_path.rglob("*")
            if path.is_file()
        }
        assert wiki_after_reprocess == wiki_before_reprocess
        assert len(provider.requests) == 3
        assert all(
            private_name not in warning
            for warning in compiled["report"]["warnings"]
            for private_name in ("Fabrizio", "Federico", "ABC Company")
        )

    asyncio.run(exercise())


@pytest.mark.parametrize("headed", (False, True), ids=("heading-free", "headed"))
def test_equivalent_txt_pdf_docx_preserve_semantics_locators_and_no_op(
    tmp_path: Path, headed: bool
) -> None:
    async def compile_format(format_name: str) -> dict[str, Any]:
        fixture, filename, content_type, ordinals = build_people_fixture(format_name, headed=headed)
        provider = ProposalProvider(
            named_subjects_proposal(("Federico La Terra", "entity", ())),
            equivalent_people_proposal(
                introduction_ordinal=ordinals[0],
                employment_ordinal=ordinals[1],
                colleague_ordinal=ordinals[2],
            ),
            equivalent_people_proposal(
                introduction_ordinal=ordinals[0],
                employment_ordinal=ordinals[1],
                colleague_ordinal=ordinals[2],
            ),
        )
        app = build_app(tmp_path / format_name, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            await upload_and_compile(
                client,
                "federico.txt",
                b"Federico La Terra is a software developer.",
                "text/plain",
            )
            uploaded, compiled = await upload_and_compile(client, filename, fixture, content_type)
            source_root = app.state.settings.source_store_path
            original_path = source_root / uploaded["original_uri"]
            original_before = original_path.read_bytes()
            session_factory = app.state.session_factory
            with session_factory() as session:
                source_record_ids_before = tuple(
                    str(record.id)
                    for record in session.scalars(
                        select(SourceRecord)
                        .where(
                            SourceRecord.document_version_id
                            == uuid.UUID(uploaded["document_version_id"])
                        )
                        .order_by(SourceRecord.ordinal)
                    )
                )
            first_graph = (await client.get("/api/v1/wiki/graph")).json()
            wiki_before = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }
            reprocessed = await client.post(
                f"/api/v1/ingestion-jobs/{uploaded['ingestion_job_id']}/reprocess"
            )
            recompiled = await client.post(
                f"/api/v1/document-versions/{uploaded['document_version_id']}/compile"
            )
            assert recompiled.json()["status"] == "COMPLETED", recompiled.json()
            graph = (await client.get("/api/v1/wiki/graph")).json()
            assert graph == first_graph

        store = WikiStore(app.state.settings.knowledge_path)
        fabrizio = store.read_page("Fabrizio Bonaccorsi")
        federico = store.read_page("Federico La Terra")
        abc = store.read_page("ABC Company")
        relationships = {
            relationship.relation: relationship for relationship in fabrizio.relationships
        }
        relation_pairs = {
            (edge["source"], edge["target"], relation)
            for edge in graph["edges"]
            for relation in edge["relations"]
        }
        wiki_after = {
            path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
            for path in app.state.settings.knowledge_path.rglob("*")
            if path.is_file()
        }
        with session_factory() as session:
            source_records = tuple(
                session.scalars(
                    select(SourceRecord)
                    .where(
                        SourceRecord.document_version_id
                        == uuid.UUID(uploaded["document_version_id"])
                    )
                    .order_by(SourceRecord.ordinal)
                )
            )
            ingestion_job = session.get(IngestionJob, uuid.UUID(uploaded["ingestion_job_id"]))
            assert ingestion_job is not None
            ingestion_timings = ingestion_job.processing_metadata["timings_ms"]

        return {
            "compiled": compiled,
            "page_ids": (fabrizio.id, federico.id, abc.id),
            "claim_ids": tuple(sorted(claim.id for claim in fabrizio.claims)),
            "relations": set(relationships),
            "employment_locators": tuple(
                item.locator for item in relationships["works_at"].provenance
            ),
            "colleague_locators": tuple(
                item.locator for item in relationships["colleague_of"].provenance
            ),
            "relation_pairs": relation_pairs,
            "original_preserved": original_path.read_bytes() == original_before == fixture,
            "source_record_ids": tuple(str(record.id) for record in source_records),
            "source_record_ids_before": source_record_ids_before,
            "wiki_no_op": wiki_after == wiki_before,
            "reprocessed": reprocessed.status_code,
            "recompiled": recompiled.json(),
            "ingestion_timings": ingestion_timings,
        }

    async def exercise() -> None:
        results = {
            format_name: await compile_format(format_name) for format_name in ("txt", "pdf", "docx")
        }
        assert len({result["page_ids"] for result in results.values()}) == 1
        assert len({result["claim_ids"] for result in results.values()}) == 1
        assert len({frozenset(result["relation_pairs"]) for result in results.values()}) == 1
        for result in results.values():
            assert result["compiled"]["status"] == "COMPLETED"
            assert result["relations"] == {"colleague_of", "works_at"}
            assert result["original_preserved"] is True
            assert result["wiki_no_op"] is True
            assert result["source_record_ids"] == result["source_record_ids_before"]
            assert result["reprocessed"] == 200
            assert result["recompiled"]["report"]["pages_created"] == []
            assert result["recompiled"]["report"]["pages_updated"] == []
            assert set(result["ingestion_timings"]) >= {
                "parse_ms",
                "normalize_ms",
                "total_ms",
            }
            assert all(value >= 0 for value in result["ingestion_timings"].values())

        docx = results["docx"]
        assert len(docx["employment_locators"]) == 2
        assert len(docx["colleague_locators"]) == 2
        assert all("paragraph" in locator for locator in docx["employment_locators"])
        assert all("paragraph" in locator for locator in docx["colleague_locators"])
        txt = results["txt"]
        pdf = results["pdf"]
        assert len(txt["employment_locators"]) == 1
        assert len(pdf["employment_locators"]) == 1

    asyncio.run(exercise())


def test_model_visible_payload_is_identity_neutral_and_timings_are_content_free(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def compile_payload(
        root: Path, filename: str, content: bytes
    ) -> tuple[str, dict[str, Any]]:
        proposal = named_subjects_proposal(("Avery Morgan", "entity", ()))
        provider = ProposalProvider(proposal)
        app = build_app(root, provider)
        rebuild_calls = 0
        lint_calls = 0
        original_rebuild = WikiStore.rebuild_index
        original_lint = WikiStore.lint

        def counted_rebuild(store: WikiStore) -> tuple[Any, ...]:
            nonlocal rebuild_calls
            rebuild_calls += 1
            return original_rebuild(store)

        def counted_lint(store: WikiStore) -> Any:
            nonlocal lint_calls
            lint_calls += 1
            return original_lint(store)

        with monkeypatch.context() as scoped:
            scoped.setattr(WikiStore, "rebuild_index", counted_rebuild)
            scoped.setattr(WikiStore, "lint", counted_lint)
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test"
            ) as client:
                _, compiled = await upload_and_compile(client, filename, content, "text/plain")
        assert rebuild_calls == 1
        assert lint_calls == 1
        request_payload = provider.requests[0].messages[1].content
        report = compiled["report"]
        assert set(report["stage_timings"]) == {
            "lock_wait_ms",
            "batch_preparation_ms",
            "provider_ms",
            "validation_grounding_relationship_ms",
            "merge_ms",
            "transactional_write_ms",
            "index_ms",
            "lint_ms",
            "manifest_and_log_ms",
            "database_commit_ms",
        }
        assert all(value >= 0 for value in report["stage_timings"].values())
        assert len(report["batches"]) == 1
        assert report["batches"][0]["attempts"][0]["attempt"] == 1
        assert filename not in json.dumps(report["stage_timings"])
        return request_payload, report

    async def exercise() -> None:
        source = "Avery Morgan leads the operations review.\n"
        first, _ = await compile_payload(tmp_path / "first", "first-name.txt", source.encode())
        second, _ = await compile_payload(
            tmp_path / "second",
            "renamed-after-delete.txt",
            b"\xef\xbb\xbf" + source.replace("\n", "\r\n").encode(),
        )
        assert first == second
        assert "source_identity" not in first
        assert "first-name.txt" not in first
        assert "renamed-after-delete.txt" not in second

    asyncio.run(exercise())


def test_adjacent_docx_window_does_not_cross_an_unrelated_paragraph(tmp_path: Path) -> None:
    async def exercise() -> None:
        document = OpenDocument()
        for paragraph in (
            "Fabrizio Bonaccorsi is an Italian web developer.",
            "The garden paragraph discusses tomato seedlings.",
            "He works at ABC Company.",
        ):
            document.add_paragraph(paragraph)
        output = BytesIO()
        document.save(output)
        proposal = named_subjects_proposal(
            ("Fabrizio Bonaccorsi", "entity", ()),
            ("ABC Company", "entity", ()),
        )
        provider = ProposalProvider(proposal)
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            _, compiled = await upload_and_compile(
                client,
                "unrelated-gap.docx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            graph = (await client.get("/api/v1/wiki/graph")).json()

        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["links_created"] == 0
        assert graph["total_edges"] == 0
        assert all(
            not page.relationships
            for _, page in WikiStore(app.state.settings.knowledge_path).page_snapshot()
        )

    asyncio.run(exercise())


@pytest.mark.parametrize(
    "sentence",
    (
        "Giuseppe La Terra and Federico La Terra are siblings.",
        (
            "Family records explicitly state that Giuseppe La Terra is the brother of "
            "Federico La Terra."
        ),
    ),
)
def test_explicit_sibling_prose_recovers_one_provenance_backed_edge_and_is_idempotent(
    tmp_path: Path, sentence: str
) -> None:
    async def exercise() -> None:
        federico = named_subjects_proposal(("Federico La Terra", "entity", ("Federico",)))
        family = named_subjects_proposal(
            ("Giuseppe La Terra", "entity", ("Giuseppe",)),
            ("Federico La Terra", "entity", ("Federico",)),
        )
        provider = ProposalProvider(federico, family, family)
        app = build_app(tmp_path, provider)
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            await upload_and_compile(
                client,
                "federico.md",
                b"Federico La Terra is a software developer.",
            )
            uploaded, compiled = await upload_and_compile(
                client,
                "family.md",
                sentence.encode(),
            )
            graph = (await client.get("/api/v1/wiki/graph")).json()
            repeated = await client.post(
                f"/api/v1/document-versions/{uploaded['document_version_id']}/compile"
            )
            wiki_before_reprocess = {
                path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
                for path in app.state.settings.knowledge_path.rglob("*")
                if path.is_file()
            }
            reprocessed = await client.post(
                f"/api/v1/ingestion-jobs/{uploaded['ingestion_job_id']}/reprocess"
            )
            recompiled = await client.post(
                f"/api/v1/document-versions/{uploaded['document_version_id']}/compile"
            )

        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["links_created"] == 1
        assert any(
            "Relationship candidate accepted: type='sibling_of'" in warning
            for warning in compiled["report"]["warnings"]
        )
        store = WikiStore(app.state.settings.knowledge_path)
        all_relationships = [
            (page.title, relationship)
            for _, page in store.page_snapshot()
            for relationship in page.relationships
        ]
        assert len(all_relationships) == 1
        source_title, relationship = all_relationships[0]
        assert source_title == "Federico La Terra"
        assert relationship.relation == "sibling_of"
        assert relationship.target_title == "Giuseppe La Terra"
        assert len(relationship.provenance) == 1
        assert relationship.provenance[0].document_version_id == uploaded["document_version_id"]
        assert graph["total_edges"] == 1
        assert "sibling_of" in graph["edges"][0]["relations"]
        assert repeated.json()["id"] == compiled["id"]
        assert reprocessed.status_code == 200
        assert recompiled.json()["status"] == "COMPLETED"
        assert recompiled.json()["report"]["pages_created"] == []
        assert recompiled.json()["report"]["pages_updated"] == []
        wiki_after_reprocess = {
            path.relative_to(app.state.settings.knowledge_path).as_posix(): path.read_bytes()
            for path in app.state.settings.knowledge_path.rglob("*")
            if path.is_file()
        }
        assert wiki_after_reprocess == wiki_before_reprocess
        assert len(provider.requests) == 3
        assert "Relationship taxonomy v1" in provider.requests[1].messages[0].content
        assert "TOPIC is a broader thematic knowledge subject" in (
            provider.requests[1].messages[0].content
        )

    asyncio.run(exercise())


def test_negated_hypothetical_quoted_ambiguous_and_cooccurring_kinship_create_no_edge(
    tmp_path: Path,
) -> None:
    unsafe_sentences = (
        "Giuseppe La Terra and Federico La Terra are not siblings.",
        "If Giuseppe La Terra and Federico La Terra were siblings, they would share a parent.",
        'The note quotes "Giuseppe La Terra and Federico La Terra are siblings."',
        "Whether Giuseppe La Terra and Federico La Terra are siblings is unknown.",
        "Giuseppe La Terra met Federico La Terra at a conference.",
    )

    async def exercise() -> None:
        proposal = named_subjects_proposal(
            ("Giuseppe La Terra", "entity", ("Giuseppe",)),
            ("Federico La Terra", "entity", ("Federico",)),
            relationships={
                "Giuseppe La Terra": [
                    {
                        "relation": "sibling_of",
                        "target_title": "Federico La Terra",
                        "source_record_ordinals": [1],
                    }
                ]
            },
        )
        provider = ProposalProvider(*(proposal for _ in unsafe_sentences))
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            for index, sentence in enumerate(unsafe_sentences, start=1):
                _, compiled = await upload_and_compile(
                    client, f"unsafe-family-{index}.md", sentence.encode()
                )
                assert compiled["status"] == "COMPLETED"
                assert compiled["report"]["links_created"] == 0
                assert any(
                    "Relationship candidate skipped: type='sibling_of'" in warning
                    for warning in compiled["report"]["warnings"]
                )
            graph = (await client.get("/api/v1/wiki/graph")).json()

        assert graph["total_edges"] == 0
        assert all(
            not page.relationships
            for _, page in WikiStore(app.state.settings.knowledge_path).page_snapshot()
        )

    asyncio.run(exercise())


@pytest.mark.parametrize(
    ("source", "subjects", "reason"),
    (
        (
            (
                "Fabrizio Bonaccorsi and Fabrizio Rossi are web developers. "
                "Fabrizio was a colleague of Federico La Terra."
            ),
            (
                ("Fabrizio Bonaccorsi", "entity", ()),
                ("Fabrizio Rossi", "entity", ()),
            ),
            "ambiguous_short_reference",
        ),
        (
            "Fabrizio Bonaccorsi met Marco Bianchi. He works at ABC Company.",
            (
                ("Fabrizio Bonaccorsi", "entity", ()),
                ("Marco Bianchi", "entity", ()),
                ("ABC Company", "entity", ()),
            ),
            "unresolved_pronoun",
        ),
    ),
)
def test_ambiguous_coreferences_emit_diagnostic_without_fabricating_edges(
    tmp_path: Path,
    source: str,
    subjects: tuple[tuple[str, str, tuple[str, ...]], ...],
    reason: str,
) -> None:
    async def exercise() -> None:
        provider = ProposalProvider(
            named_subjects_proposal(("Federico La Terra", "entity", ())),
            named_subjects_proposal(*subjects),
        )
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            await upload_and_compile(
                client,
                "federico.md",
                b"Federico La Terra is a software developer.",
            )
            _, compiled = await upload_and_compile(client, "ambiguous.md", source.encode())
            graph = (await client.get("/api/v1/wiki/graph")).json()

        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["links_created"] == 0
        assert any(f"reason='{reason}'" in warning for warning in compiled["report"]["warnings"])
        assert graph["total_edges"] == 0
        assert all(
            not page.relationships
            for _, page in WikiStore(app.state.settings.knowledge_path).page_snapshot()
        )

    asyncio.run(exercise())


def test_safelist_recovers_directional_relations_resolves_aliases_and_skips_unresolved(
    tmp_path: Path,
) -> None:
    subjects = (
        ("Avery Morgan", "entity", ("Avery",)),
        ("Acme Corporation", "entity", ("ACME",)),
        ("University of Malta", "entity", ()),
        ("Project Atlas", "entity", ("Atlas",)),
        ("Python", "concept", ()),
        ("Valletta", "entity", ()),
        ("Unrelated Garden", "topic", ()),
    )
    relation_subjects = (
        ("Avery Morgan", "entity", ("Avery",)),
        ("ACME", "entity", ("Acme Corporation",)),
        ("University of Malta", "entity", ()),
        ("Project Atlas", "entity", ("Atlas",)),
        ("Python", "concept", ()),
        ("Valletta", "entity", ()),
    )
    explicit = (
        "Avery Morgan founded ACME. Avery Morgan owns ACME. Avery Morgan works at ACME. "
        "Avery Morgan studied at University of Malta. "
        "Project Atlas was created by Avery Morgan. Avery Morgan uses Python. "
        "Avery Morgan lives in Valletta."
    )
    model_relationships: dict[str, list[dict[str, object]]] = {
        "Avery Morgan": [
            {
                "relation": "works_for",
                "target_title": "ACME",
                "source_record_ordinals": [1],
            },
            {
                "relation": "works_at",
                "target_title": "Missing Organization",
                "source_record_ordinals": [1],
            },
        ]
    }

    async def exercise() -> None:
        provider = ProposalProvider(
            named_subjects_proposal(*subjects),
            named_subjects_proposal(*relation_subjects, relationships=model_relationships),
        )
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            await upload_and_compile(
                client,
                "subjects.md",
                (
                    b"Avery Morgan, Acme Corporation (ACME), University of Malta, Project Atlas, "
                    b"Python, Valletta, and Unrelated Garden are named subjects."
                ),
            )
            garden_path = next(
                path
                for path in app.state.settings.knowledge_path.rglob("*.md")
                if path.stem == "unrelated-garden"
            )
            garden_before = garden_path.read_bytes()
            _, compiled = await upload_and_compile(
                client, "explicit-relations.md", explicit.encode()
            )
            graph = (await client.get("/api/v1/wiki/graph")).json()

        assert compiled["status"] == "COMPLETED"
        assert compiled["report"]["links_created"] == 7
        assert any(
            "reason='unresolved_endpoint'" in warning for warning in compiled["report"]["warnings"]
        )
        store = WikiStore(app.state.settings.knowledge_path)
        relations = {
            (page.title, relationship.relation, relationship.target_title)
            for _, page in store.page_snapshot()
            for relationship in page.relationships
        }
        assert relations == {
            ("Avery Morgan", "founder_of", "Acme Corporation"),
            ("Avery Morgan", "lives_in", "Valletta"),
            ("Avery Morgan", "owner_of", "Acme Corporation"),
            ("Avery Morgan", "studied_at", "University of Malta"),
            ("Avery Morgan", "uses_technology", "Python"),
            ("Avery Morgan", "works_at", "Acme Corporation"),
            ("Project Atlas", "created_by", "Avery Morgan"),
        }
        assert graph["total_edges"] == 5
        assert garden_path.read_bytes() == garden_before

    asyncio.run(exercise())


def test_page_kind_guidance_preserves_distinct_entity_concept_and_topic_outputs(
    tmp_path: Path,
) -> None:
    async def exercise() -> None:
        proposal = named_subjects_proposal(
            ("Mira Chen", "entity", ()),
            ("Least Privilege", "concept", ()),
            ("Information Security", "topic", ()),
        )
        provider = ProposalProvider(proposal)
        app = build_app(tmp_path, provider)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            _, compiled = await upload_and_compile(
                client,
                "security.md",
                (
                    b"Mira Chen applies Least Privilege while documenting the broader "
                    b"Information Security program."
                ),
            )
            pages = (await client.get("/api/v1/wiki/pages")).json()

        assert compiled["status"] == "COMPLETED"
        assert {(page["title"], page["type"]) for page in pages} == {
            ("Mira Chen", "entity"),
            ("Least Privilege", "concept"),
            ("Information Security", "topic"),
        }
        assert not list((app.state.settings.knowledge_path / "wiki" / "datasets").glob("*.md"))

    asyncio.run(exercise())
